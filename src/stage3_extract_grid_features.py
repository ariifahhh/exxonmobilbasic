"""Stream reservoir grids into compact, reproducible case-level ML features."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
FEATURE_COLUMNS = [
    "Pore Volume (PORV)", "Porosity (PORO)",
    "Permeability I (PERMX)", "Permeability J (PERMY)", "Permeability K (PERMZ)",
    "Net-to-Gross Ratio (NTG)",
    "Transmissibility I (TRANX)", "Transmissibility J (TRANY)", "Transmissibility K (TRANZ)",
    "Pressure (PRESSURE)", "Water Saturation (SWAT)", "Oil Saturation (SOIL)",
    "Gas Saturation (SGAS)", "Solution Gas-Oil Ratio (RS)", "Vapor Oil-Gas Ratio (RV)",
]


def split_for_case(case_id: int) -> str:
    if 1 <= case_id <= 70:
        return "train"
    if 71 <= case_id <= 85:
        return "validation"
    if 86 <= case_id <= 100:
        return "test"
    raise ValueError(f"Unexpected case_id: {case_id}")


def case_id_from_sheet(sheet_name: str) -> int:
    match = re.search(r"(\d+)", sheet_name)
    if not match:
        raise ValueError(f"Cannot identify case number from sheet {sheet_name!r}")
    return int(match.group(1))


def normalise_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def summarise_sheet(source_path: Path, sheet_name: str) -> tuple[dict[str, object], dict[str, object]]:
    """Read one 90k-row sheet once, retaining only online summary statistics."""
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [normalise_header(value) for value in next(rows)]
    positions = {header: index for index, header in enumerate(headers)}
    missing = [column for column in FEATURE_COLUMNS if column not in positions]
    if missing:
        workbook.close()
        raise ValueError(f"{source_path.name} / {sheet_name} is missing columns: {missing}")

    count = 0
    sums = np.zeros(len(FEATURE_COLUMNS), dtype=float)
    sums_sq = np.zeros(len(FEATURE_COLUMNS), dtype=float)
    minima = np.full(len(FEATURE_COLUMNS), np.inf)
    maxima = np.full(len(FEATURE_COLUMNS), -np.inf)
    missing_values = 0
    indices = [positions[column] for column in FEATURE_COLUMNS]

    for row in rows:
        values = np.array([row[index] for index in indices], dtype=float)
        finite = np.isfinite(values)
        missing_values += int((~finite).sum())
        if not finite.all():
            values = np.where(finite, values, 0.0)
        sums += values
        sums_sq += values * values
        minima = np.minimum(minima, values)
        maxima = np.maximum(maxima, values)
        count += 1
    workbook.close()

    means = sums / count
    variances = np.maximum(sums_sq / count - means * means, 0.0)
    feature_row: dict[str, object] = {
        "case_id": case_id_from_sheet(sheet_name),
        "split": split_for_case(case_id_from_sheet(sheet_name)),
        "source_file": source_path.name,
        "source_sheet": sheet_name,
    }
    for index, column in enumerate(FEATURE_COLUMNS):
        short_name = re.search(r"\(([^)]+)\)", column).group(1).lower()
        feature_row.update({
            f"{short_name}_mean": means[index],
            f"{short_name}_std": float(np.sqrt(variances[index])),
            f"{short_name}_min": minima[index],
            f"{short_name}_max": maxima[index],
        })
    audit_row = {
        "case_id": feature_row["case_id"],
        "split": feature_row["split"],
        "source_file": source_path.name,
        "source_sheet": sheet_name,
        "grid_rows": count,
        "feature_columns": len(FEATURE_COLUMNS),
        "missing_feature_values": missing_values,
        "status": "pass" if count == 90365 and missing_values == 0 else "review",
    }
    return feature_row, audit_row


def run(input_dir: Path, processed_dir: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Extract every grid's features and join the four provided uncertainty values."""
    processed_dir.mkdir(parents=True, exist_ok=True)
    feature_path = processed_dir / "grid_case_features.csv"
    audit_path = processed_dir / "grid_feature_audit.csv"
    existing = pd.read_csv(feature_path) if feature_path.exists() else pd.DataFrame()
    existing_ids = set(existing["case_id"].astype(int)) if not existing.empty else set()
    feature_rows = existing.to_dict("records") if not existing.empty else []
    audit_rows = pd.read_csv(audit_path).to_dict("records") if audit_path.exists() else []

    source_files = sorted(
        path for path in input_dir.glob("*.xlsx")
        if path.name.startswith(("03 ", "04 ", "05 ", "06 ", "07 ", "08 "))
    )
    for source_path in source_files:
        workbook = load_workbook(source_path, read_only=True, data_only=True)
        sheet_names = workbook.sheetnames
        workbook.close()
        for sheet_name in sheet_names:
            case_id = case_id_from_sheet(sheet_name)
            if case_id in existing_ids:
                continue
            feature_row, audit_row = summarise_sheet(source_path, sheet_name)
            feature_rows.append(feature_row)
            audit_rows.append(audit_row)
            pd.DataFrame(feature_rows).sort_values("case_id").to_csv(feature_path, index=False)
            pd.DataFrame(audit_rows).sort_values("case_id").to_csv(audit_path, index=False)
            print(f"Completed Case {case_id:03d} ({audit_row['split']})")

    features = pd.DataFrame(feature_rows).sort_values("case_id").drop_duplicates("case_id", keep="last")
    audit = pd.DataFrame(audit_rows).sort_values("case_id").drop_duplicates("case_id", keep="last")
    uncertainty = pd.read_csv(processed_dir / "case_uncertainties.csv")
    features = features.merge(uncertainty, on="case_id", how="left", validate="one_to_one")
    if features["case_id"].nunique() != 100:
        raise ValueError(f"Expected 100 cases, found {features['case_id'].nunique()}")
    if features.isna().any().any():
        missing = features.columns[features.isna().any()].tolist()
        raise ValueError(f"Features contain missing values: {missing}")
    features.to_csv(feature_path, index=False)
    audit.to_csv(audit_path, index=False)
    return features, audit


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract compact ML features from reservoir grid workbooks.")
    parser.add_argument("--input-dir", type=Path, default=PROJECT_ROOT / "TestCases")
    parser.add_argument("--processed-dir", type=Path, default=PROJECT_ROOT / "data" / "processed")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    features, audit = run(args.input_dir, args.processed_dir)
    print(f"Completed {len(features)} cases with {len(features.columns) - 4} model features.")
    print(audit.groupby(["split", "status"]).size().to_string())
